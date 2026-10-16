from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "Gebruiksaanwijzing_HealthNoord.docx"
LOGO_PATH = ROOT / "assets" / "hoofdlogo-healthNoord.png"
SCREENSHOT_DIR = ROOT / "docs" / "handleiding"
# Figure label -> screenshot file in SCREENSHOT_DIR (all made with synthetic data).
SCREENSHOTS = {
    "Hoofdscherm": "hoofdscherm.png",
    "Bestandskeuze en voorcontrole": "bestandskeuze_voorcontrole.png",
    "Filtergroepen en lijsten": "filtergroepen_lijsten.png",
    "Voortgang tijdens verwerking": "voortgang.png",
    "Resultatenvenster en log": "resultatenvenster_log.png",
}

NAVY = "1F2D54"
ORANGE = "EA590C"
PALE_BLUE = "EEF2F8"
PALE_ORANGE = "FFF2E9"
MID_GREY = "637083"
LINE_GREY = "C9D1DC"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_border(cell, color: str = LINE_GREY, size: str = "8") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Pagina ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 29, NAVY),
        ("Heading 1", 20, NAVY),
        ("Heading 2", 14, ORANGE),
        ("Heading 3", 11, NAVY),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name in {"Title", "Heading 1"} else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "HEALTH NOORD  |  ANONIMISATIETOOL"
    header.style = styles["Caption"]
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    add_page_number(section.footer.paragraphs[0])


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)


def add_number(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Number")


def add_callout(document: Document, title: str, text: str, warning: bool = False) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, PALE_ORANGE if warning else PALE_BLUE)
    set_cell_border(cell, ORANGE if warning else LINE_GREY)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    heading = paragraph.add_run(f"{title}  ")
    heading.bold = True
    heading.font.color.rgb = RGBColor.from_string(ORANGE if warning else NAVY)
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_screenshot(document: Document, label: str, instructions: str) -> None:
    """Insert the screenshot for `label` from SCREENSHOT_DIR, or a placeholder if it is missing."""
    image_path = SCREENSHOT_DIR / SCREENSHOTS.get(label, "")
    if image_path.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Inches(6.5))
    else:
        add_screenshot_placeholder(document, label, instructions)
        return
    caption = document.add_paragraph(f"Figuur: {label}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.style = document.styles["Caption"]
    caption.runs[0].italic = True
    caption.runs[0].font.color.rgb = RGBColor.from_string(MID_GREY)


def add_screenshot_placeholder(document: Document, label: str, instructions: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.85)
    cell = table.cell(0, 0)
    cell.width = Inches(6.85)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "F7F8FA")
    set_cell_border(cell, ORANGE, "12")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(29)
    paragraph.paragraph_format.space_after = Pt(5)
    marker = paragraph.add_run(f"[SCREENSHOT {label.upper()} HIER INVOEGEN]")
    marker.bold = True
    marker.font.size = Pt(13)
    marker.font.color.rgb = RGBColor.from_string(ORANGE)
    detail = cell.add_paragraph(instructions)
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail.paragraph_format.space_after = Pt(29)
    detail.runs[0].font.size = Pt(9)
    detail.runs[0].font.color.rgb = RGBColor.from_string(MID_GREY)
    caption = document.add_paragraph(f"Figuur: {label}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.style = document.styles["Caption"]
    caption.runs[0].italic = True
    caption.runs[0].font.color.rgb = RGBColor.from_string(MID_GREY)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_values in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row_values):
            cells[index].text = text
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def add_title_page(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(40)
    if LOGO_PATH.exists():
        paragraph.add_run().add_picture(str(LOGO_PATH), width=Inches(6.2))

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(16)
    run = subtitle.add_run("Gebruikershandleiding")
    run.bold = True
    run.font.size = Pt(27)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    product = document.add_paragraph()
    product.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = product.add_run("Health Noord Anonimisatie-tool")
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(ORANGE)

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_before = Pt(18)
    intro.add_run("Handleiding voor de Windows-applicatie (.exe)\n")
    intro.add_run("Lokale verwerking van Nederlandstalige zorgteksten")

    add_callout(
        document,
        "Belangrijk",
        "Controleer de uitvoer altijd voordat je deze deelt. Een anonimisatietool kan herkenbare gegevens missen.",
        warning=True,
    )
    version = document.add_paragraph("Versie van deze handleiding: 16 oktober 2026")
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.runs[0].font.size = Pt(9)
    version.runs[0].font.color.rgb = RGBColor.from_string(MID_GREY)
    document.add_page_break()


def build_manual() -> Path:
    document = Document()
    configure_document(document)
    add_title_page(document)

    document.add_heading("Inhoud", level=1)
    for item in [
        "1. Over deze tool",
        "2. Voorbereiding en bestanden",
        "3. De interface",
        "4. Een document anonimiseren",
        "5. Voortgang en annuleren",
        "6. Resultaten en logbestand",
        "7. Extra lijsten en whitelist",
        "8. Instellingen onthouden",
        "9. Privacy en veilig gebruik",
        "10. Problemen oplossen",
        "11. Controlelijst voor vrijgave",
    ]:
        add_bullet(document, item)
    document.add_page_break()

    document.add_heading("1. Over deze tool", level=1)
    document.add_paragraph(
        "Health Noord Anonimisatie-tool is een Windows-desktopprogramma voor het lokaal anonimiseren van tekstbestanden. "
        "De tool gebruikt Deduce voor entiteitsherkenning en kan optionele eigen lijsten toepassen op namen en locaties. "
        "Een whitelist kan geselecteerde medische of domeintermen bewust ongewijzigd laten."
    )
    add_callout(
        document,
        "Geen garantie",
        "Automatische herkenning kan gegevens missen of onbedoeld vervangen. Beoordeel het resultaat altijd inhoudelijk en volg het privacy- en informatiebeveiligingsbeleid van je organisatie.",
        warning=True,
    )
    document.add_heading("Wat de tool wel en niet doet", level=2)
    add_bullet(document, "Verwerkt één UTF-8 of ander ondersteund tekstbestand per run.")
    add_bullet(document, "Schrijft een nieuw geanonimiseerd tekstbestand en een apart logbestand.")
    add_bullet(document, "Werkt lokaal; de ingevoerde tekst wordt niet naar een clouddienst verstuurd.")
    add_bullet(document, "Vervangt geen menselijke privacycontrole en is geen garantie op volledige anonimisering.")

    document.add_heading("2. Voorbereiding en bestanden", level=1)
    document.add_heading("Benodigdheden", level=2)
    add_bullet(document, "Een Windows 10- of Windows 11-computer.")
    add_bullet(document, "De door je organisatie verstrekte HealthNoordAnonymizer.exe.")
    add_bullet(document, "Een invoerbestand in .txt-formaat. UTF-8 is aanbevolen.")
    add_bullet(document, "Optioneel: extra lijsten en/of een whitelist, elk als .txt-bestand.")

    document.add_heading("Bestandstypen", level=2)
    add_table(
        document,
        ["Bestand", "Inhoud", "Formaat"],
        [
            ["Invoer", "Tekst die verwerkt moet worden", "Eén .txt-document"],
            ["Patiëntenlijst", "Patiëntnamen die zeker vervangen moeten worden", "Eén naam per regel"],
            ["Medewerkerslijst", "Medewerkersnamen die zeker vervangen moeten worden", "Eén naam per regel"],
            ["Locatielijst", "Locaties of instellingen die extra herkend moeten worden", "Eén term per regel"],
            ["Whitelist", "Termen die bewust ongewijzigd blijven", "Eén woord of woordgroep per regel"],
        ],
    )
    document.add_paragraph(
        "Lege regels in lijsten worden genegeerd. Vergelijking is hoofdletterongevoelig. Bewaar lijsten op een locatie met passende toegangsbeveiliging."
    )

    document.add_heading("3. De interface", level=1)
    document.add_paragraph(
        "De hoofdschermindeling bestaat uit een bestandsgedeelte, filterinstellingen, runbediening en een resultaatvoorbeeld. "
        "Bestandsdialogen openen telkens in de map waar de applicatie staat."
    )
    add_screenshot(
        document,
        "Hoofdscherm",
        "Voeg hier een afbeelding in van het volledige hoofdvenster. Houd bestandsnamen en voorbeeldtekst niet-persoonlijk.",
    )
    document.add_heading("Onderdelen", level=2)
    add_table(
        document,
        ["Onderdeel", "Gebruik"],
        [
            ["Bestanden", "Kies invoerbestand, uitvoermap, optionele logmap, lijsten en whitelist."],
            ["Filter op", "Schakel Deduce-filtergroepen aan of uit voor deze run."],
            ["Anonimiseer bestand", "Voert eerst een controle en bevestiging uit, daarna start de verwerking."],
            ["Annuleer verwerking", "Vraagt bevestiging en stopt na afronding van de huidige tekstchunk."],
            ["Voortgang", "Toont fase, percentage, gelezen bytes en na enige tijd een resttijdschatting."],
            ["Resultaatvoorbeeld", "Toont een begrensd voorbeeld van de uitvoer; het volledige bestand staat op schijf."],
        ],
    )

    document.add_heading("4. Een document anonimiseren", level=1)
    for step in [
        "Start HealthNoordAnonymizer.exe. Wacht tot de initialisatie gereed is.",
        "Kies bij Invoer .txt het tekstbestand dat je in deze run wilt verwerken. Het invoerbestand wordt niet onthouden.",
        "Kies bij Uitvoermap waar het geanonimiseerde bestand moet worden opgeslagen.",
        "Kies desgewenst een aparte Logmap. Zonder aparte keuze wordt het log in de uitvoermap geplaatst.",
        "Kies eventueel patiënten-, medewerkers- en locatielijsten. Elke lijst kan afzonderlijk worden weggelaten.",
        "Kies eventueel een whitelist. Gebruik deze alleen voor termen die ongewijzigd in de uitvoer mogen voorkomen.",
        "Controleer de filtergroepen. Uitgeschakelde groepen worden niet door Deduce verwerkt. Is 'Namen' uitgevinkt, dan worden ook de patiënten- en medewerkerslijst niet gebruikt; is 'Locaties' uitgevinkt, dan ook de locatielijst niet.",
        "Klik op Anonimiseer bestand. Lees de voorcontrole en controleer bestandsnaam, bestandsgrootte en uitvoerlocaties.",
        "Bevestig de start. Tijdens de verwerking worden bestandskeuzes en filters vergrendeld om wijzigingen midden in een run te voorkomen.",
    ]:
        add_number(document, step)
    add_screenshot(
        document,
        "Bestandskeuze en voorcontrole",
        "Voeg hier een screenshot toe met voorbeeldbestanden en de bevestiging voor de start. Gebruik uitsluitend synthetische namen.",
    )
    add_screenshot(
        document,
        "Filtergroepen en lijsten",
        "Voeg hier een screenshot toe van de filtergroepcheckboxes en geselecteerde lijstsoorten; scherm persoonlijke bestandspaden af.",
    )

    document.add_heading("Voorcontrole", level=2)
    document.add_paragraph(
        "Voor de run controleert de tool of het invoerbestand bestaat en leesbaar is, of de optionele .txt-bestanden beschikbaar zijn en of de uitvoer- en logmappen beschrijfbaar zijn. "
        "Bij een probleem verschijnt een korte lijst met punten die eerst moeten worden hersteld."
    )
    add_callout(
        document,
        "Bestandsgrootte",
        "Bij grote bestanden kan de verwerking langer duren. De bevestiging toont de grootte voordat de run start.",
    )

    document.add_heading("5. Voortgang en annuleren", level=1)
    document.add_paragraph(
        "Tijdens de run toont de statusregel de actuele fase, zoals tekst verwerken of uitvoer opslaan. De voortgangsbalk is gebaseerd op de invoerbytes. "
        "Na voldoende meettijd verschijnt een schatting van de resterende tijd; dit blijft een benadering, omdat tekstinhoud de herkenningstijd beïnvloedt."
    )
    add_screenshot(
        document,
        "Voortgang tijdens verwerking",
        "Voeg hier een screenshot toe van een actieve run met fase, percentage en resttijd. Gebruik een niet-gevoelig voorbeeld.",
    )
    document.add_heading("Annuleren", level=2)
    add_number(document, "Klik op Annuleer verwerking.")
    add_number(document, "Bevestig de annuleervraag.")
    add_number(
        document,
        "De huidige tekstchunk wordt eerst afgerond. Daarna stopt de run, worden tijdelijke uitvoerbestanden verwijderd en wordt een log met status Geannuleerd opgeslagen.",
    )
    add_callout(
        document,
        "Niet geforceerd afsluiten",
        "Sluit het proces niet via Taakbeheer tijdens verwerking. Gebruik de annuleerknop zodat tijdelijke bestanden netjes worden opgeruimd.",
        warning=True,
    )

    document.add_heading("6. Resultaten en logbestand", level=1)
    document.add_heading("Uitvoer", level=2)
    document.add_paragraph(
        "Na een geslaagde run wordt een nieuw bestand met de invoernaam en een tijdstempel opgeslagen. De naam bevat doorgaans de aanduiding `_anonymized_`. "
        "Een bestaand bestand wordt niet overschreven. Het resultaatvoorbeeld in het venster is begrensd; het volledige resultaat staat in het opgeslagen bestand."
    )
    document.add_heading("Resultatenvenster", level=2)
    document.add_paragraph(
        "Na succes verschijnt automatisch een scrollbaar venster met de volledige logresultaten, waaronder per filtergroep of die aan of uit stond. "
        "Onderin open je met Open resultaat het bestand, met Open uitvoermap de map en met Open log het logbestand. Sluit het met Sluiten, Escape of het kruisje. "
        "Daarna blijven dezelfde acties beschikbaar als kleine links (Resultaat · Uitvoermap · Log) rechts op de statusregel, tot je een nieuwe run start."
    )
    add_screenshot(
        document,
        "Resultatenvenster en log",
        "Voeg hier een screenshot toe van het resultatenvenster met synthetische metadata en tellingen. Verwijder eventuele gebruikersnamen of gevoelige paden.",
    )
    document.add_heading("Wat staat in het log?", level=2)
    add_table(
        document,
        ["Logveld", "Betekenis"],
        [
            ["Status", "Geslaagd, Geannuleerd of Mislukt."],
            ["Starttijd en duur", "Wanneer de run begon en hoeveel seconden deze duurde."],
            ["Bestandsmetadata", "Invoerbestandsnaam, grootte en bij succes uitvoernaam; geen volledige paden."],
            ["Instellingen", "Actieve en uitgeschakelde filtergroepen en alleen de ja/nee-status voor whitelist en extra lijsten."],
            ["Vervangingstotalen", "Deduce-vervangingen, extra-lijstvervangingen en hun som."],
            ["Per type", "Aantallen per herkenningstype, bijvoorbeeld BSN, persoon of extra locatie."],
            ["Foutdetails", "Bij mislukking een foutcategorie en korte, waar mogelijk padgemaskeerde melding."],
        ],
    )
    add_callout(
        document,
        "Geannuleerd of mislukt",
        "Ook dan wordt een statuslog bewaard. Er wordt geen gedeeltelijk geanonimiseerd eindbestand gepubliceerd.",
    )

    document.add_heading("7. Extra lijsten en whitelist", level=1)
    document.add_heading("Extra lijsten", level=2)
    document.add_paragraph(
        "Extra lijsten bevatten namen en locaties die zeker vervangen moeten worden. Ze worden vóór Deduce toegepast, zodat een naam uit de lijst in zijn geheel wordt vervangen, "
        "inclusief tussenvoegsels en dubbele achternaam; daarna anonimiseert Deduce de rest van de tekst. Namenlijsten bevatten volledige namen (voornaam en achternaam); "
        "locatielijsten bevatten één locatie per regel. Een voor- of achternaam die meerdere mensen in de lijst delen, wordt ook los vervangen, met een eigen pseudoniem. "
        "Alle lijsten worden in één keer doorlopen en de langste treffer wint, zodat de straat 'Harmen Systrasingel' uit de locatielijst voorgaat op de voornaam 'Harmen'. "
        "Bij een straat uit de locatielijst wordt het huisnummer meegenomen, ook met toevoeging ('12a', '258-K'). Een locatie wordt alleen als hele regel herkend; "
        "zet varianten daarom als aparte regels in de lijst. Een postcode op 'MG' in hoofdletters ('9257MG'), die Deduce overslaat, wordt als locatie vervangen; "
        "'1000 mg' in kleine letters blijft staan als dosering. "
        "De tool kan personen pseudonimiseren met consistente placeholders en telt extra vervangingen afzonderlijk in het log."
    )
    document.add_paragraph(
        "Hoofdletters, spaties, regelafbrekingen, koppeltekens en apostroffen mogen afwijken van de lijst. Een dubbele achternaam zoals "
        "'van den Broek - Zoon' blijft één geheel. Een losse voor- of achternaam, of een naam van één woord (eventueel met tussenvoegsel), "
        "wordt alleen vervangen als die in de tekst met een hoofdletter staat: 'Arts' en 'de Goede' wel, de gewone woorden 'arts' en 'de goede' niet. "
        "Volledige namen van meer woorden worden altijd vervangen. Een achternaam of naam van één woord die ook een gewoon Nederlands woord is (zoals 'Linker' of 'Anders') "
        "wordt aan het begin van een zin niet vervangen, omdat de hoofdletter daar van het zinsbegin komt; na een titel of voorletter ('mevr. Anders') en midden in een zin wel. "
        "Voornamen uit een volledige naam (zoals 'Wil' uit 'Wil Bakker') worden ook aan het begin van een zin vervangen."
    )
    document.add_heading("Whitelist", level=2)
    document.add_paragraph(
        "Een whitelistbestand bevat één woord of woordgroep per regel. Overeenkomsten zijn hoofdletterongevoelig en volgen woordgrenzen. Een gevonden term wordt vóór herkenning beschermd en daarna in de oorspronkelijke schrijfwijze teruggezet. "
        "Staat een term ook in de patiënten-, medewerkers- of locatielijst, dan gaat die lijst voor en wordt de term toch geanonimiseerd. "
        "Uitzondering: een term van één woord die ook een gewoon Nederlands woord is (zoals 'Kamer' of 'Zondag') blijft staan als die op de whitelist staat."
    )
    add_callout(
        document,
        "Bewuste uitzondering",
        "Whitelist-termen blijven zichtbaar in de uitvoer. Plaats daarom alleen vaste medische of domeintermen op de whitelist, nooit namen of andere gegevens die geanonimiseerd moeten worden.",
        warning=True,
    )
    document.add_paragraph(
        "Voorbeeld: als `Gabriele-de Vries syndroom` een vaste aandoeningsnaam is die behouden mag blijven, zet dan de volledige woordgroep op één regel. Voeg niet alleen `Gabriele` of `de Vries` toe als die woorden ook als persoonsnaam voorkomen."
    )

    document.add_heading("8. Instellingen onthouden", level=1)
    document.add_paragraph(
        "De tool bewaart de filterkeuzes. De andere keuzes worden bewust nooit onthouden; selecteer dit voor iedere run opnieuw."
    )

    document.add_heading("9. Privacy en veilig gebruik", level=1)
    add_bullet(document, "De verwerking gebeurt lokaal op de computer; tekst wordt niet naar een cloudservice gestuurd.")
    add_bullet(document, "Logs bevatten geen tekstfragmenten, whitelisttermen, lijstinhoud of volledige bestandspaden.")
    add_bullet(document, "Bestandsnamen en vervangingstellingen kunnen nog steeds operationele informatie onthullen; deel logs daarom alleen met bevoegde personen.")
    add_bullet(document, "Bewaar input, uitvoer, lijsten en whitelist volgens het bewaarbeleid van je organisatie.")
    add_bullet(document, "Controleer de uitvoer op resterende directe en indirecte identificatoren voordat je deze deelt.")
    add_bullet(document, "Gebruik een whitelist uitsluitend als de behouden term aantoonbaar in de gedeelde tekst mag staan.")
    add_callout(
        document,
        "Verantwoordelijkheid",
        "De gebruiker blijft verantwoordelijk voor controle van de uitvoer, juiste selectie van lijsten en naleving van AVG- en organisatiebeleid.",
        warning=True,
    )

    document.add_heading("10. Problemen oplossen", level=1)
    add_table(
        document,
        ["Melding of situatie", "Mogelijke oorzaak en oplossing"],
        [
            ["Invoerbestand ontbreekt", "Selecteer een bestaand .txt-bestand. Het invoerbestand wordt niet automatisch hersteld uit een vorige sessie."],
            ["Alleen .txt-bestanden worden ondersteund", "Exporteer of sla de brontekst eerst op als .txt. CSV, PDF en Word-bestanden worden niet direct verwerkt."],
            ["Uitvoer- of logmap niet beschrijfbaar", "Kies een map waarin je account bestanden mag aanmaken, bijvoorbeeld een persoonlijke werkmap."],
            ["Optionele lijst is niet leesbaar", "Controleer of het bestand bestaat, een .txt-extensie heeft en niet door een ander programma is vergrendeld."],
            ["Verwerking lijkt lang te duren", "Bekijk fase, bytevoortgang en resttijd. Grote chunks en complexe tekst kunnen meer rekentijd vragen."],
            ["Annuleren duurt nog even", "De huidige chunk moet eerst klaar zijn. Daarna wordt de run veilig afgebroken en verschijnt een annuleerlog."],
            ["Whitelistterm is vervangen", "Controleer de exacte woordgroep, spelling en woordgrenzen in het whitelistbestand. Gebruik geen te brede korte termen. Staat de term ook in een extra lijst, dan wordt hij bewust toch geanonimiseerd, behalve als het één gewoon Nederlands woord is."],
            ["Gewoon woord wordt als naam vervangen", "Het woord staat in de namenlijst en midden in een zin met hoofdletter. Zet het op de whitelist of haal het uit de namenlijst."],
            ["Windows SmartScreen waarschuwt", "Gebruik alleen de door je organisatie verstrekte executable en volg het interne softwarebeleid."],
        ],
    )

    document.add_heading("11. Controlelijst voor vrijgave", level=1)
    for item in [
        "Is het juiste invoerbestand geselecteerd en is de naam/grootte in de voorcontrole correct?",
        "Zijn alleen noodzakelijke extra lijsten geselecteerd?",
        "Zijn whitelisttermen gecontroleerd op ongewenst behoud van persoons- of identificatiegegevens?",
        "Zijn uitvoer- en logmap geschikt en toegankelijk voor de bedoelde gebruiker?",
        "Is de verwerking volledig afgerond en is de status Geslaagd?",
        "Is het volledige uitvoerbestand gecontroleerd voordat het wordt gedeeld?",
        "Worden uitvoer en log veilig bewaard of verwijderd volgens organisatiebeleid?",
    ]:
        add_bullet(document, item)

    document.add_paragraph()
    add_callout(
        document,
        "Bij twijfel",
        "Deel de uitvoer niet totdat de privacyverantwoordelijke of afgesproken reviewer de tekst heeft gecontroleerd.",
        warning=True,
    )

    document.core_properties.title = "Gebruikershandleiding Health Noord Anonimisatie-tool"
    document.core_properties.subject = "Uitgebreide handleiding voor Windows .exe"
    document.core_properties.author = "Health Noord"
    document.core_properties.keywords = "Health Noord, anonimisatie, handleiding, Windows"
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_manual())